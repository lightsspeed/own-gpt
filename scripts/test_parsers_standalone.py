import sys, re, csv, json, os
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ = {}  # dummy

# Test parser functions directly
def parse_pdf(path):
    try:
        import fitz
        docs = []
        with fitz.open(path) as doc:
            for page_num, page in enumerate(doc, start=1):
                text = page.get_text().strip()
                if text:
                    docs.append({"page": page_num, "content": text[:100]})
        return docs
    except Exception as e:
        return [{"error": str(e)}]

def parse_docx(path):
    try:
        from docx import Document
        doc = Document(str(path))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        return [{"paragraphs": len(paras), "sample": paras[0][:100]}] if paras else []
    except Exception as e:
        return [{"error": str(e)}]

def parse_pptx(path):
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        slides = len(prs.slides)
        return [{"slides": slides}]
    except Exception as e:
        return [{"error": str(e)}]

def parse_csv_simple(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        dialect = csv.Sniffer().sniff(f.read(8192))
        f.seek(0)
        reader = csv.reader(f, dialect)
        headers = next(reader, None)
        rows = list(reader)
    return [{"headers": headers, "rows": len(rows)}]

def parse_xlsx(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        info = {}
        for name in wb.sheetnames:
            ws = wb[name]
            rows = sum(1 for _ in ws.iter_rows(values_only=True))
            info[name] = rows
        wb.close()
        return [{"sheets": info}]
    except Exception as e:
        return [{"error": str(e)}]

def parse_html_simple(path):
    try:
        from bs4 import BeautifulSoup
        with open(path, encoding="utf-8", errors="replace") as f:
            soup = BeautifulSoup(f.read(), "html.parser")
        title = soup.find("title")
        title_text = title.get_text(strip=True) if title else "none"
        body_text = soup.get_text(separator=" ", strip=True)[:100]
        return [{"title": title_text, "body": body_text}]
    except Exception as e:
        return [{"error": str(e)}]

base = Path("test_ingest_src")
for p in sorted(base.rglob("*")):
    if not p.is_file():
        continue
    ext = p.suffix.lower()
    parser = {
        ".pdf": parse_pdf,
        ".docx": parse_docx,
        ".pptx": parse_pptx,
        ".csv": parse_csv_simple,
        ".xlsx": parse_xlsx,
        ".html": parse_html_simple,
        ".htm": parse_html_simple,
        ".txt": lambda p: [{"content": p.read_text(encoding="utf-8")[:100]}],
    }.get(ext)
    if parser:
        result = parser(p)
        print("[{}] {} → {}".format(ext, p.name, json.dumps(result, ensure_ascii=False)[:120]))
